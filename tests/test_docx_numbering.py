"""T10.53 — Word 다단계 자동 번호("3.2" 등) 계산 테스트.

워드가 문단 텍스트가 아니라 numPr(직접 또는 스타일 상속)로 번호를 그려주면
python-docx가 읽는 텍스트에는 그 번호가 없다 — 실측(exdoc docx 10개, 2026-09-01):
제목 108개 중 95개(88%)가 이 패턴. 실제 문서 문구는 쓰지 않는다(T10.52의 교훈) —
가상 예시로 재현한다.
"""

from __future__ import annotations

from lxml import etree

from parser.utils.numbering import NumberingResolver, _parse_levels

_WNS = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'


def _lvl_xml(ilvl: int, num_fmt: str, lvl_text: str, start: int = 1) -> str:
    return (
        f'<w:lvl {_WNS} w:ilvl="{ilvl}">'
        f'<w:start w:val="{start}"/>'
        f'<w:numFmt w:val="{num_fmt}"/>'
        f'<w:lvlText w:val="{lvl_text}"/>'
        f"</w:lvl>"
    )


def _numbering_xml(levels: list[tuple]) -> bytes:
    lvl_xml = "".join(_lvl_xml(*lvl) for lvl in levels)
    return (
        f'<w:numbering {_WNS}>'
        f'<w:abstractNum w:abstractNumId="0">{lvl_xml}</w:abstractNum>'
        f'<w:num w:numId="1"><w:abstractNumId w:val="0"/></w:num>'
        f"</w:numbering>"
    ).encode("utf-8")


def _add_multilevel_numbering(document, num_id: str, abstract_id: str, levels: list[tuple]) -> None:
    """(ilvl, numFmt, lvlText[, start]) 목록으로 다단계 번호 정의를 문서에 추가한다."""
    lvl_xml = "".join(_lvl_xml(*lvl) for lvl in levels)
    abstract_xml = f'<w:abstractNum {_WNS} w:abstractNumId="{abstract_id}">{lvl_xml}</w:abstractNum>'
    num_xml = f'<w:num {_WNS} w:numId="{num_id}"><w:abstractNumId w:val="{abstract_id}"/></w:num>'

    numbering_root = document.part.numbering_part.element
    numbering_root.append(etree.fromstring(abstract_xml))
    numbering_root.append(etree.fromstring(num_xml))


def _set_style_numpr(document, style_name: str, num_id: str, ilvl: int) -> None:
    """스타일의 `pPr`에 numPr을 건다 — 그 스타일을 쓰는 모든 문단이 상속받는다.

    `document.styles[...]`는 스타일 **이름**("Heading 1")으로 찾는다 — 저장된
    XML의 `styleId`("Heading1")와 다르다.
    """
    style = document.styles[style_name]
    ppr = style.element.get_or_add_pPr()
    numpr_xml = f'<w:numPr {_WNS}><w:ilvl w:val="{ilvl}"/><w:numId w:val="{num_id}"/></w:numPr>'
    ppr.append(etree.fromstring(numpr_xml))


def _set_direct_numpr(paragraph, num_id: str, ilvl: int) -> None:
    """문단 자신의 `pPr`에 numPr을 건다 — 스타일 상속보다 우선해야 한다."""
    ppr = paragraph._p.get_or_add_pPr()
    numpr_xml = f'<w:numPr {_WNS}><w:ilvl w:val="{ilvl}"/><w:numId w:val="{num_id}"/></w:numPr>'
    ppr.append(etree.fromstring(numpr_xml))


class _FakeStyle:
    def __init__(self, style_id: str) -> None:
        self.style_id = style_id


class _FakeParagraph:
    """`NumberingResolver.number_for()`가 쓰는 속성만 흉내 낸 가짜 문단."""

    def __init__(self, style_id: str) -> None:
        self.style = _FakeStyle(style_id)
        self._p = etree.fromstring(f"<w:p {_WNS}/>")


class TestNumberingResolverUnit:
    """`NumberingResolver`만 단독으로 — 실제 문서 문구는 쓰지 않는다."""

    def test_two_level_decimal_sequence(self):
        """1. / 1.1 / 1.2 / 2. 순서 — 상위 레벨이 바뀌면 하위가 리셋된다."""
        levels = _parse_levels(_numbering_xml([(0, "decimal", "%1."), (1, "decimal", "%1.%2")]))
        # H1(ilvl0)·H2(ilvl1) 스타일을 각각 등록해 실제 장/절 구조를 흉내 낸다.
        resolver = NumberingResolver({"H1": ("1", 0), "H2": ("1", 1)}, levels)

        assert resolver.number_for(_FakeParagraph("H1")) == "1."
        assert resolver.number_for(_FakeParagraph("H2")) == "1.1"
        assert resolver.number_for(_FakeParagraph("H2")) == "1.2"
        assert resolver.number_for(_FakeParagraph("H1")) == "2."
        # 상위(H1)가 다시 늘었으니 하위는 처음부터 다시 시작한다.
        assert resolver.number_for(_FakeParagraph("H2")) == "2.1"

    def test_missing_numbering_definition_is_a_safe_noop(self):
        """번호 매기기 정의가 아예 없는 문서 — 항상 빈 문자열."""
        resolver = NumberingResolver({}, {})
        assert resolver.number_for(_FakeParagraph("H1")) == ""

    def test_unsupported_format_is_skipped(self):
        """bullet 같은 형식은 억지로 숫자를 만들지 않고 번호 없음으로 본다."""
        levels = _parse_levels(_numbering_xml([(0, "bullet", "")]))
        resolver = NumberingResolver({"H1": ("1", 0)}, levels)
        assert resolver.number_for(_FakeParagraph("H1")) == ""

    def test_roman_and_letter_formats(self):
        """decimal 외 지원 형식도 렌더링된다."""
        levels = _parse_levels(
            _numbering_xml([(0, "upperRoman", "%1."), (1, "lowerLetter", "%2)")])
        )
        resolver = NumberingResolver({"H1": ("1", 0), "H2": ("1", 1)}, levels)

        assert resolver.number_for(_FakeParagraph("H1")) == "I."
        assert resolver.number_for(_FakeParagraph("H2")) == "a)"
        assert resolver.number_for(_FakeParagraph("H2")) == "b)"


class TestDocxParserNumberingIntegration:
    """`DocxParser`가 실제 docx(python-docx로 조립)에 적용했을 때."""

    def _build(self, tmp_path, *, direct_override=False):
        import docx

        source = docx.Document()
        source.add_paragraph("표지")
        source.add_paragraph("전문 구조 정의", style="Heading 1")
        source.add_paragraph("본문 1")
        source.add_paragraph("데이터 형식", style="Heading 2")
        source.add_paragraph("본문 1.1")
        source.add_paragraph("공통 구조", style="Heading 2")
        source.add_paragraph("본문 1.2")
        last_h1 = source.add_paragraph("코드 정의", style="Heading 1")
        source.add_paragraph("본문 2")

        _add_multilevel_numbering(
            source, num_id="200", abstract_id="200",
            levels=[(0, "decimal", "%1."), (1, "decimal", "%1.%2 ")],
        )
        _set_style_numpr(source, "Heading 1", "200", 0)
        _set_style_numpr(source, "Heading 2", "200", 1)

        if direct_override:
            # 문단 하나만 다른 번호 매기기 계열을 직접 지정 — 스타일 상속보다
            # 우선해야 한다.
            _add_multilevel_numbering(
                source, num_id="201", abstract_id="201",
                levels=[(0, "upperRoman", "%1.")],
            )
            _set_direct_numpr(last_h1, "201", 0)

        path = tmp_path / "sample.docx"
        source.save(path)
        return path

    def test_numbers_are_prepended_to_headings(self, tmp_path):
        from parser.formats.docx_parser import DocxParser

        path = self._build(tmp_path)
        document = DocxParser(asset_dir=tmp_path / "assets").parse(path)
        headings = [c.heading for c in document.chunks if c.heading]

        assert headings == ["1. 전문 구조 정의", "1.1 데이터 형식", "1.2 공통 구조", "2. 코드 정의"]

    def test_numbered_heading_is_searchable_via_body_content(self, tmp_path):
        """새 컬럼 없이 기존 FTS 경로로 검색되는가 — 번호가 본문 content에도 남는다."""
        from parser.formats.docx_parser import DocxParser

        path = self._build(tmp_path)
        document = DocxParser(asset_dir=tmp_path / "assets").parse(path)
        body = "\n".join(c.content for c in document.chunks if c.type.value == "text")

        assert "1.1 데이터 형식" in body

    def test_direct_paragraph_numpr_overrides_style(self, tmp_path):
        """문단 자신의 numPr이 스타일 상속보다 우선한다(워드의 실제 우선순위)."""
        from parser.formats.docx_parser import DocxParser

        path = self._build(tmp_path, direct_override=True)
        document = DocxParser(asset_dir=tmp_path / "assets").parse(path)
        headings = [c.heading for c in document.chunks if c.heading]

        # 마지막 Heading 1("코드 정의")만 로마 숫자 계열로 직접 지정했다.
        assert headings[-1] == "I. 코드 정의"

    def test_no_numbering_definition_falls_back_to_plain_heading(self, tmp_path):
        """번호 매기기 정의가 없는 순정 docx는 지금까지와 동일하게 동작한다."""
        import docx

        from parser.formats.docx_parser import DocxParser

        source = docx.Document()
        source.add_paragraph("통신 메시지 구조 정의", style="Heading 1")
        source.add_paragraph("본문")
        path = tmp_path / "plain.docx"
        source.save(path)

        document = DocxParser(asset_dir=tmp_path / "assets").parse(path)
        headings = [c.heading for c in document.chunks if c.heading]
        assert headings == ["통신 메시지 구조 정의"]
